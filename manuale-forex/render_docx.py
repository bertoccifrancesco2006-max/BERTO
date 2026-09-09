# -*- coding: utf-8 -*-
import re, sys
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from parser import load_all, BOX_KINDS

SERIF, SANS = 'Liberation Serif', 'Liberation Sans'
BOXFILL = {'flow': 'F4F4F4', 'def': 'EEF1F5', 'err': 'F5F0EE', 'rule': 'EEF3EF', 'note': 'F4F2EC', 'key': 'EFF0F2', 'warn': 'F2F2F2'}
BOXBAR = {'flow': '8F8F8F', 'def': '2F3A46', 'err': '5B3A30', 'rule': '2F4A38', 'note': '5A5030', 'key': '3A3A44', 'warn': '444444'}


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), v)
    return e


def set_cell_bg(cell, hexcol):
    cell._tc.get_or_add_tcPr().append(el('w:shd', val='clear', color='auto', fill=hexcol))


def cell_borders(cell, left=None, box=None):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        e = el('w:' + side, val='single', sz='4', space='0', color='A8A8A8')
        if side == 'left' and left:
            e = el('w:left', val='single', sz='24', space='0', color=left)
        b.append(e)
    tcPr.append(b)


def add_field(par, instr):
    r = par.add_run()
    f1 = el('w:fldChar', fldCharType='begin'); r._r.append(f1)
    r2 = par.add_run()
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    r2._r.append(it)
    r3 = par.add_run()
    r3._r.append(el('w:fldChar', fldCharType='separate'))
    r4 = par.add_run('1')
    r5 = par.add_run()
    r5._r.append(el('w:fldChar', fldCharType='end'))
    return r4


def runs_from_md(par, text, base_font=SERIF, size=10.5, color=None):
    text = text.replace('->', '→').replace('=>', '⇒')
    parts = re.split(r'(\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*))', text)
    for p in parts:
        if not p:
            continue
        bold = ital = False
        if p.startswith('**') and p.endswith('**'):
            p = p[2:-2]; bold = True
        elif p.startswith('*') and p.endswith('*'):
            p = p[1:-1]; ital = True
        r = par.add_run(p)
        r.font.name = base_font
        r._element.rPr.rFonts.set(qn('w:eastAsia'), base_font)
        r.font.size = Pt(size)
        r.bold = bold; r.italic = ital
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


class DocxRenderer:
    def __init__(self, doc):
        self.d = doc
        self.pending_cols = None

    def para(self, container, text, style=None, font=SERIF, size=10.5, bold=False,
             align=None, space_after=4, space_before=0, indent=0, bullet=None, color=None, keep=False):
        p = container.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(space_after); pf.space_before = Pt(space_before)
        pf.line_spacing = 1.13
        if align:
            p.alignment = align
        if indent:
            pf.left_indent = Mm(indent)
            pf.first_line_indent = Mm(-4)
        if keep:
            pf.keep_with_next = True
        pf.widow_control = True
        if bullet:
            r = p.add_run(bullet + '  ')
            r.font.name = SERIF; r.font.size = Pt(size)
        runs_from_md(p, text, font, size, color)
        if bold:
            for r in p.runs:
                r.bold = True
        return p

    def render(self, toks, container=None, inbox=False):
        c = container if container is not None else self.d
        for t in toks:
            k = t[0]
            if k == 'p':
                self.para(c, t[1], size=9.7 if inbox else 10.5,
                          align=WD_ALIGN_PARAGRAPH.JUSTIFY if not inbox else WD_ALIGN_PARAGRAPH.LEFT,
                          space_after=3 if inbox else 5)
            elif k == 'h':
                self.heading(c, t[1], t[2])
            elif k == 'ul':
                for sub, txt in t[1]:
                    self.para(c, txt, size=9.7 if inbox else 10.5, indent=8 if sub else 4,
                              bullet='–' if sub else '•', space_after=2)
            elif k == 'ol':
                for n, txt in t[1]:
                    self.para(c, txt, size=9.7 if inbox else 10.5, indent=5, bullet=n + '.', space_after=2)
            elif k == 'cl':
                for txt in t[1]:
                    self.para(c, txt, size=9.7 if inbox else 10.5, indent=6,
                              bullet='\u2610', space_after=6)
            elif k == 'table':
                self.table(c, t[1])
            elif k == 'box':
                self.box(c, t[1], t[2], t[3])
            elif k == 'cmd':
                self.cmd(c, t[1], t[2])

    def heading(self, c, lvl, txt):
        if lvl == 1 and len(self.d.paragraphs) > 1:
            pb = self.d.add_paragraph()
            pb.add_run().add_break(WD_BREAK.PAGE)
            pb.paragraph_format.space_after = Pt(0)
        style = {1: 'Heading 1', 2: 'Heading 2', 3: 'Heading 3', 4: 'Heading 4'}[lvl]
        p = c.add_paragraph(style=style)
        sizes = {1: 20, 2: 14.5, 3: 11.5, 4: 10.5}
        runs_from_md(p, txt, SANS, sizes[lvl])
        for r in p.runs:
            r.bold = True
            r.italic = (lvl == 4)
            r.font.color.rgb = RGBColor.from_string('1A1A1A' if lvl <= 2 else '333333')
        pf = p.paragraph_format
        pf.keep_with_next = True
        pf.space_before = Pt({1: 4, 2: 16, 3: 9, 4: 6}[lvl])
        pf.space_after = Pt({1: 6, 2: 6, 3: 3, 4: 2}[lvl])
        if lvl <= 2:
            pPr = p._p.get_or_add_pPr()
            b = OxmlElement('w:pBdr')
            b.append(el('w:bottom', val='single', sz='12' if lvl == 1 else '6', space='4', color='1A1A1A' if lvl == 1 else '9A9A9A'))
            pPr.append(b)

    def table(self, c, rows):
        if not rows:
            return
        ncol = max(len(r) for r in rows)
        t = self.d.add_table(rows=0, cols=ncol)
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        if self.pending_cols and len(self.pending_cols) == ncol:
            tot = sum(self.pending_cols)
            widths = [Mm(170 * x / tot) for x in self.pending_cols]
            self.pending_cols = None
        else:
            widths = [Mm(170 / ncol)] * ncol
        for ri, r in enumerate(rows):
            r = r + [''] * (ncol - len(r))
            row = t.add_row()
            for ci, txt in enumerate(r):
                cell = row.cells[ci]
                cell.width = widths[ci]
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                runs_from_md(p, txt, SANS if ri == 0 else SERIF, 9)
                if ri == 0:
                    for run in p.runs:
                        run.bold = True
                    set_cell_bg(cell, 'E4E4E4')
                elif ri % 2 == 0:
                    set_cell_bg(cell, 'F7F7F7')
        self.d.add_paragraph().paragraph_format.space_after = Pt(3)

    def box(self, c, kind, title, body):
        if kind == 'flow':
            return self.flowbox(title, body)
        label, _ = BOX_KINDS[kind]
        head = label if not title else '%s · %s' % (label, title)
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        cell = t.rows[0].cells[0]
        cell.width = Mm(170)
        set_cell_bg(cell, BOXFILL[kind])
        cell_borders(cell, left=BOXBAR[kind])
        p0 = cell.paragraphs[0]
        p0.paragraph_format.space_after = Pt(3)
        runs_from_md(p0, head, SANS, 8.6)
        for r in p0.runs:
            r.bold = True
        self.render(body, container=cell, inbox=True)
        cell.paragraphs[-1].paragraph_format.space_after = Pt(0)
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)

    def flowbox(self, title, body):
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        cell = t.rows[0].cells[0]
        cell.width = Mm(170)
        set_cell_bg(cell, 'F4F4F4')
        cell_borders(cell, left='8F8F8F')
        steps = [x[1] for x in body if x[0] == 'p']
        first = True
        for i, st in enumerate(steps):
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            if '|' in st:
                a, b = st.split('|', 1)
                runs_from_md(p, a.strip(), SANS, 9.2)
                for r in p.runs:
                    r.bold = True
                br = p.add_run(); br.add_break()
                r2 = p.add_run(b.strip())
                r2.font.name = SANS; r2.font.size = Pt(7.8); r2.italic = True
                r2.font.color.rgb = RGBColor.from_string('5A5A5A')
            else:
                runs_from_md(p, st, SANS, 9.2)
                for r in p.runs:
                    r.bold = True
            if i < len(steps) - 1:
                q = cell.add_paragraph()
                q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                q.paragraph_format.space_after = Pt(0)
                q.paragraph_format.space_before = Pt(0)
                rr = q.add_run('\u2193')
                rr.font.name = SANS; rr.font.size = Pt(8.5)
                rr.font.color.rgb = RGBColor.from_string('5A5A5A')
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)

    def notes(self, n):
        p = self.d.add_paragraph()
        r = p.add_run('APPUNTI PERSONALI')
        r.font.name = SANS; r.font.size = Pt(8); r.bold = True
        r.font.color.rgb = RGBColor.from_string('5A5A5A')
        p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(4)
        for i in range(n):
            q = self.d.add_paragraph()
            q.paragraph_format.space_after = Pt(10)
            q.add_run(' ')
            pPr = q._p.get_or_add_pPr()
            b = OxmlElement('w:pBdr')
            b.append(el('w:bottom', val='single', sz='4', space='6', color='C9C9C9'))
            pPr.append(b)

    def cmd(self, c, name, arg):
        if name == 'pagebreak':
            p = self.d.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
        elif name == 'notes':
            self.notes(int(arg or 6))
        elif name == 'cols':
            self.pending_cols = [float(x) for x in (arg or '').split(',') if x.strip()]
        elif name == 'space':
            p = self.d.add_paragraph(); p.paragraph_format.space_after = Pt(float(arg or 5) * 2)
        elif name == 'hr':
            p = self.d.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            b = OxmlElement('w:pBdr'); b.append(el('w:bottom', val='single', sz='6', space='2', color='9A9A9A'))
            pPr.append(b)
        elif name == 'toc':
            p = self.d.add_paragraph()
            add_field(p, ' TOC \\o "1-2" \\h \\z \\u ')
            self.para(self.d, '*Se l\'indice appare vuoto: cliccalo e premi F9 (Word) oppure Strumenti → Aggiorna → Indici (LibreOffice) per generarlo con i numeri di pagina.*',
                      size=8.6, font=SANS, color='5A5A5A')
        elif name == 'cover':
            self.cover()

    def cover(self):
        d = self.d
        for _ in range(4):
            d.add_paragraph()
        self.para(d, 'MANUALE DI STUDIO DEL FOREX', font=SANS, size=30, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
        self.para(d, 'Dalle basi alla lettura del mercato', font=SANS, size=14,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
        self.para(d, 'Guida didattica costruita sul percorso video', font=SERIF, size=12,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        self.para(d, '**«Da Base a Pro» (E1–E24) di Giuliano**', font=SANS, size=12.5,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
        self.para(d, '25 lezioni · novembre–dicembre 2024', font=SERIF, size=11,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
        for _ in range(10):
            d.add_paragraph()
        self.para(d, '*Documento di studio a uso personale. Non è consulenza finanziaria né un invito a investire. Il trading comporta il rischio di perdere il capitale.*',
                  size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color='5A5A5A')
        p = d.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)


def build(srcdir, outfile):
    toks = load_all(srcdir)
    d = Document()
    st = d.settings.element
    st.append(el('w:updateFields', val='true'))
    sec = d.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.left_margin = Mm(22); sec.right_margin = Mm(18)
    sec.top_margin = Mm(20); sec.bottom_margin = Mm(18)
    sec.header_distance = Mm(11); sec.footer_distance = Mm(10)
    sec.different_first_page_header_footer = True

    n = d.styles['Normal']
    n.font.name = SERIF; n.font.size = Pt(10.5)
    n._element.rPr.rFonts.set(qn('w:eastAsia'), SERIF)
    n.paragraph_format.space_after = Pt(5)

    # intestazione
    hp = sec.header.paragraphs[0]
    hp.text = ''
    r = hp.add_run('MANUALE DI STUDIO DEL FOREX · Da Base a Pro')
    r.font.name = SANS; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string('5A5A5A')
    hp.add_run('\t\t')
    r2 = hp.add_run('E1–E24 · Giuliano')
    r2.font.name = SANS; r2.font.size = Pt(8); r2.font.color.rgb = RGBColor.from_string('5A5A5A')
    pPr = hp._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr'); b.append(el('w:bottom', val='single', sz='4', space='3', color='9A9A9A'))
    pPr.append(b)

    fp = sec.footer.paragraphs[0]
    fp.text = ''
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = add_field(fp, ' PAGE ')
    rr.font.name = SANS; rr.font.size = Pt(9); rr.bold = True
    pPr = fp._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr'); b.append(el('w:top', val='single', sz='4', space='6', color='9A9A9A'))
    pPr.append(b)

    r = DocxRenderer(d)
    r.render(toks)
    d.save(outfile)
    return outfile


if __name__ == '__main__':
    print('DOCX:', build(sys.argv[1], sys.argv[2]))
